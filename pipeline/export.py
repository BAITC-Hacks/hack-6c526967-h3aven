"""
Модуль экспорта протокола совещания в DOCX.

Создает профессиональный документ:
- Шапка (название, дата, участники)
- Повестка дня
- Полный транскрипт с разделением по спикерам
- Ключевые решения
- Таблица поручений (ответственный, срок, приоритет)
- Саммари
"""

import os
from html import escape
from pathlib import Path
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# Цвета для спикеров
SPEAKER_COLORS = [
    RGBColor(0x1A, 0x73, 0xE8),  # Синий
    RGBColor(0xE8, 0x45, 0x3C),  # Красный
    RGBColor(0x0B, 0x8E, 0x43),  # Зеленый
    RGBColor(0xF9, 0xAB, 0x00),  # Желтый
    RGBColor(0x9C, 0x27, 0xB0),  # Фиолетовый
    RGBColor(0x00, 0xBC, 0xD4),  # Бирюзовый
    RGBColor(0xFF, 0x57, 0x22),  # Оранжевый
    RGBColor(0x60, 0x7D, 0x8B),  # Серо-синий
]

PRIORITY_LABELS = {
    "high": "Высокий",
    "medium": "Средний",
    "low": "Низкий",
}


def export_to_docx(
    transcript_data: dict,
    analysis_data: dict = None,
    output_dir: str = "temp",
    filename: str = None,
) -> str:
    """
    Создает профессиональный DOCX-протокол совещания.

    Args:
        transcript_data: результат транскрибации (segments, duration, etc.).
        analysis_data: результат LLM-анализа (summary, tasks, participants).
        output_dir: папка для сохранения.
        filename: имя файла (без расширения).

    Returns:
        Путь к созданному DOCX-файлу.
    """
    os.makedirs(output_dir, exist_ok=True)

    if filename is None:
        source = transcript_data.get("source_file", "meeting")
        name = os.path.splitext(source)[0]
        filename = f"Протокол_{name}"

    output_path = os.path.join(output_dir, f"{filename}.docx")

    document = Document()

    # Настраиваем стили
    style = document.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)

    # Поля документа
    for section in document.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    # ──────────────────────────────────────────────
    # ШАПКА
    # ──────────────────────────────────────────────
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("ПРОТОКОЛ СОВЕЩАНИЯ")
    title_run.bold = True
    title_run.font.size = Pt(16)

    if analysis_data and analysis_data.get("title"):
        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle.add_run(analysis_data["title"])
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Метаданные
    meta_para = document.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    date_str = datetime.now().strftime("%d.%m.%Y")
    if analysis_data and analysis_data.get("date_mentioned"):
        date_str = analysis_data["date_mentioned"]

    duration_min = int(transcript_data.get("duration", 0)) // 60
    duration_sec = int(transcript_data.get("duration", 0)) % 60

    meta_run = meta_para.add_run(
        f"Дата: {date_str}\n"
        f"Длительность: {duration_min} мин {duration_sec} сек\n"
        f"Количество участников: {transcript_data.get('num_speakers', 0)}\n"
        f"Язык: {transcript_data.get('language', 'ru').upper()}"
    )
    meta_run.font.size = Pt(11)
    meta_run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    document.add_paragraph()  # Отступ

    # ──────────────────────────────────────────────
    # УЧАСТНИКИ
    # ──────────────────────────────────────────────
    if analysis_data and analysis_data.get("participants"):
        _add_heading(document, "Участники совещания")
        for i, p in enumerate(analysis_data["participants"], 1):
            name = p.get("name", "")
            role = p.get("role", "")
            text = f"{i}. {name}"
            if role:
                text += f" -- {role}"
            document.add_paragraph(text)
        document.add_paragraph()

    # ──────────────────────────────────────────────
    # ПОВЕСТКА ДНЯ
    # ──────────────────────────────────────────────
    if analysis_data and analysis_data.get("agenda"):
        _add_heading(document, "Повестка дня")
        for i, item in enumerate(analysis_data["agenda"], 1):
            document.add_paragraph(f"{i}. {item}")
        document.add_paragraph()

    # ──────────────────────────────────────────────
    # САММАРИ
    # ──────────────────────────────────────────────
    if analysis_data and analysis_data.get("summary"):
        _add_heading(document, "Краткое содержание (Саммари)")
        para = document.add_paragraph(analysis_data["summary"])
        para.paragraph_format.first_line_indent = Cm(1.25)
        document.add_paragraph()

    # ──────────────────────────────────────────────
    # КЛЮЧЕВЫЕ РЕШЕНИЯ
    # ──────────────────────────────────────────────
    if analysis_data and analysis_data.get("key_decisions"):
        _add_heading(document, "Принятые решения")
        for i, decision in enumerate(analysis_data["key_decisions"], 1):
            document.add_paragraph(f"{i}. {decision}")
        document.add_paragraph()

    # ──────────────────────────────────────────────
    # ПОРУЧЕНИЯ
    # ──────────────────────────────────────────────
    tasks = []
    if analysis_data:
        tasks = analysis_data.get("tasks", [])

    if tasks:
        _add_heading(document, "Поручения")

        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Заголовки
        headers = ["No", "Поручение", "Ответственный", "Срок", "Приоритет"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(10)

        # Данные
        for i, task in enumerate(tasks, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = task.get("task", "")
            row[2].text = task.get("assignee", "Не указан")
            row[3].text = task.get("deadline", "Не указан")
            priority = task.get("priority", "medium")
            row[4].text = PRIORITY_LABELS.get(priority, priority)

            for cell in row:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

        # Ширина колонок
        widths = [Cm(1), Cm(7), Cm(3.5), Cm(3), Cm(2)]
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]

        document.add_paragraph()

    # ──────────────────────────────────────────────
    # ТРАНСКРИПТ
    # ──────────────────────────────────────────────
    segments = transcript_data.get("segments", [])
    if segments:
        _add_heading(document, "Полный транскрипт совещания")

        # Маппинг спикеров для цветов
        speakers = list(dict.fromkeys(s["speaker"] for s in segments))
        speaker_color_map = {}
        for i, sp in enumerate(speakers):
            speaker_color_map[sp] = SPEAKER_COLORS[i % len(SPEAKER_COLORS)]

        for seg in segments:
            speaker = seg.get("speaker", "SPEAKER_00")
            start = seg.get("start", 0)
            end = seg.get("end", 0)
            text = seg.get("text", "")

            start_m, start_s = divmod(int(start), 60)
            end_m, end_s = divmod(int(end), 60)

            para = document.add_paragraph()

            # Таймкод
            time_run = para.add_run(f"[{start_m:02d}:{start_s:02d} - {end_m:02d}:{end_s:02d}] ")
            time_run.font.size = Pt(9)
            time_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

            # Спикер
            speaker_run = para.add_run(f"{speaker}: ")
            speaker_run.bold = True
            speaker_run.font.size = Pt(11)
            speaker_run.font.color.rgb = speaker_color_map.get(
                speaker, RGBColor(0, 0, 0)
            )

            # Текст
            text_run = para.add_run(text)
            text_run.font.size = Pt(11)

    # ──────────────────────────────────────────────
    # ПОДПИСЬ
    # ──────────────────────────────────────────────
    document.add_paragraph()
    footer = document.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(
        f"Протокол сформирован автоматически системой SÖZDE\n"
        f"{datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    footer_run.italic = True

    document.save(output_path)
    print(f"[export] DOCX сохранен: {output_path}")
    return output_path


def _add_heading(document, text: str):
    """Добавляет форматированный заголовок раздела."""
    heading = document.add_heading(text, level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x73, 0xE8)


def _pdf_fonts():
    """Register Unicode fonts available on Windows and common Linux images."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
         Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    ]
    for regular, bold in candidates:
        if regular.is_file():
            pdfmetrics.registerFont(TTFont("Protocol", str(regular)))
            pdfmetrics.registerFont(TTFont("Protocol-Bold", str(bold if bold.is_file() else regular)))
            return "Protocol", "Protocol-Bold"
    return "Helvetica", "Helvetica-Bold"


def export_to_pdf(
    transcript_data: dict,
    analysis_data: dict = None,
    output_dir: str = "temp",
    filename: str = None,
) -> str:
    """Create a page-numbered Unicode PDF protocol."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    analysis_data = analysis_data or {}
    destination = Path(output_dir).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    if filename is None:
        source = Path(transcript_data.get("source_file", "meeting")).stem
        filename = f"Протокол_{source}"
    output_path = destination / f"{filename}.pdf"
    regular_font, bold_font = _pdf_fonts()

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ProtocolBody", parent=styles["BodyText"], fontName=regular_font,
        fontSize=9.5, leading=13, textColor=colors.HexColor("#24313d"),
        spaceAfter=5,
    )
    heading = ParagraphStyle(
        "ProtocolHeading", parent=body, fontName=bold_font, fontSize=13,
        leading=16, textColor=colors.HexColor("#174f45"), spaceBefore=12,
        spaceAfter=7,
    )
    title_style = ParagraphStyle(
        "ProtocolTitle", parent=heading, fontSize=18, leading=22,
        alignment=TA_CENTER, textColor=colors.HexColor("#17324d"), spaceAfter=10,
    )
    small = ParagraphStyle(
        "ProtocolSmall", parent=body, fontSize=8, leading=10,
        textColor=colors.HexColor("#5d6873"),
    )

    def page_footer(canvas, document):
        canvas.saveState()
        canvas.setFont(regular_font, 8)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(1.6 * cm, 1.0 * cm, "SÖZDE | локальная обработка")
        canvas.drawRightString(19.4 * cm, 1.0 * cm, f"Страница {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        str(output_path), pagesize=A4, rightMargin=1.6 * cm, leftMargin=1.6 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title=analysis_data.get("title", "Протокол совещания"),
        author="SÖZDE",
    )
    story = [Paragraph("ПРОТОКОЛ СОВЕЩАНИЯ", title_style)]
    if analysis_data.get("title"):
        story.append(Paragraph(escape(str(analysis_data["title"])), heading))

    duration = int(transcript_data.get("duration", 0))
    metadata = (
        f"Дата формирования: {datetime.now():%d.%m.%Y %H:%M}<br/>"
        f"Источник: {escape(str(transcript_data.get('source_file', '')))}<br/>"
        f"Длительность: {duration // 60} мин {duration % 60} сек | "
        f"Спикеров: {transcript_data.get('num_speakers', 0)} | "
        f"Язык: {escape(str(transcript_data.get('language', 'unknown')).upper())}"
    )
    story.extend([Paragraph(metadata, small), Spacer(1, 8)])

    sections = [
        ("Саммари", analysis_data.get("summary")),
        ("Повестка", analysis_data.get("agenda")),
        ("Принятые решения", analysis_data.get("key_decisions")),
        ("Риски", analysis_data.get("risks")),
    ]
    for section_title, content in sections:
        if not content:
            continue
        story.append(Paragraph(section_title, heading))
        if isinstance(content, list):
            for index, item in enumerate(content, 1):
                story.append(Paragraph(f"{index}. {escape(str(item))}", body))
        else:
            story.append(Paragraph(escape(str(content)).replace("\n", "<br/>"), body))

    tasks = analysis_data.get("tasks", [])
    if tasks:
        story.append(Paragraph("Поручения", heading))
        rows = [[
            Paragraph("No", small), Paragraph("Поручение", small),
            Paragraph("Ответственный", small), Paragraph("Срок", small),
            Paragraph("Статус", small),
        ]]
        for index, task in enumerate(tasks, 1):
            rows.append([
                Paragraph(str(index), small),
                Paragraph(escape(str(task.get("task", ""))), small),
                Paragraph(escape(str(task.get("assignee", "Не указан"))), small),
                Paragraph(escape(str(task.get("deadline", "Не указан"))), small),
                Paragraph(escape(str(task.get("status", "pending"))), small),
            ])
        table = Table(rows, colWidths=[0.7 * cm, 7.6 * cm, 3.4 * cm, 2.8 * cm, 2.2 * cm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17324d")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), bold_font),
            ("FONTNAME", (0, 1), (-1, -1), regular_font),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5df")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7f9")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend([table, Spacer(1, 8)])

    segments = transcript_data.get("segments", [])
    if segments:
        story.extend([PageBreak(), Paragraph("Полный транскрипт", heading)])
        for segment in segments:
            start = int(segment.get("start", 0))
            end = int(segment.get("end", 0))
            label = escape(str(segment.get("speaker", "SPEAKER_00")))
            text = escape(str(segment.get("text", "")))
            line = (
                f"<font color='#6b7280'>[{start // 60:02d}:{start % 60:02d} - "
                f"{end // 60:02d}:{end % 60:02d}]</font> "
                f"<b>{label}:</b> {text}"
            )
            story.append(Paragraph(line, body))

    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    print(f"[export] PDF сохранен: {output_path}")
    return str(output_path)


def export_to_markdown(
    transcript_data: dict,
    analysis_data: dict = None,
    output_dir: str = "temp",
    filename: str = None,
) -> str:
    """Create a portable Markdown protocol for local archives and SED import."""
    analysis_data = analysis_data or {}
    destination = Path(output_dir).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    source = Path(transcript_data.get("source_file", "meeting")).stem
    output_path = destination / f"{filename or f'Протокол_{source}'}.md"
    lines = [f"# {analysis_data.get('title', 'Протокол совещания')}", ""]
    if analysis_data.get("summary"):
        lines.extend(["## Саммари", "", str(analysis_data["summary"]), ""])
    if analysis_data.get("key_decisions"):
        lines.extend(["## Решения", ""])
        lines.extend(f"- {item}" for item in analysis_data["key_decisions"])
        lines.append("")
    if analysis_data.get("tasks"):
        lines.extend([
            "## Поручения", "",
            "| Поручение | Ответственный | Срок | Приоритет | Статус |",
            "|---|---|---|---|---|",
        ])
        for task in analysis_data["tasks"]:
            values = [
                task.get("task", ""), task.get("assignee", "Не указан"),
                task.get("deadline", "Не указан"), task.get("priority", "medium"),
                task.get("status", "pending"),
            ]
            lines.append("| " + " | ".join(str(value).replace("|", "\\|") for value in values) + " |")
        lines.append("")
    lines.extend(["## Транскрипт", ""])
    for segment in transcript_data.get("segments", []):
        seconds = int(segment.get("start", 0))
        lines.append(
            f"**[{seconds // 60:02d}:{seconds % 60:02d}] "
            f"{segment.get('speaker', 'SPEAKER_00')}:** {segment.get('text', '')}"
        )
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return str(output_path)


def export_all(transcript_data: dict, output_dir: str = "temp") -> dict:
    """Generate all judge-facing document formats in one call."""
    analysis = transcript_data.get("analysis") or {}
    return {
        "docx_path": export_to_docx(transcript_data, analysis, output_dir),
        "pdf_path": export_to_pdf(transcript_data, analysis, output_dir),
        "markdown_path": export_to_markdown(transcript_data, analysis, output_dir),
    }
