from io import BytesIO
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from docx import Document
from docx.shared import Pt


def _time(seconds: float) -> str:
    return f"{int(seconds // 60):02}:{int(seconds % 60):02}"


def protocol_text(meeting) -> str:
    people = ", ".join(meeting.speaker_mapping.values()) or "Не указаны"
    lines = ["ПРОТОКОЛ СОВЕЩАНИЯ", f"Дата: {meeting.meeting_date}", f"Участники: {people}", "", "КРАТКОЕ СОДЕРЖАНИЕ", meeting.summary, "", "ПОРУЧЕНИЯ"]
    for index, task in enumerate(meeting.tasks, 1):
        lines.append(f"{index}. {task.task} — {task.assignee or 'Требует уточнения'} — {task.deadline or 'Срок не указан'}")
    lines.extend(["", "ПОЛНЫЙ ТРАНСКРИПТ"])
    for segment in (meeting.transcript or {}).get("segments", []):
        speaker = meeting.speaker_mapping.get(segment["speaker"], segment["speaker"])
        lines.append(f"[{_time(segment['start'])}] {speaker}: {segment['text']}")
    return "\n".join(lines)


def docx_bytes(meeting) -> bytes:
    document = Document()
    document.add_heading("ПРОТОКОЛ СОВЕЩАНИЯ", 0)
    document.add_paragraph(f"Дата: {meeting.meeting_date}")
    document.add_paragraph(f"Участники: {', '.join(meeting.speaker_mapping.values()) or 'Не указаны'}")
    document.add_heading("Краткое содержание", 1)
    document.add_paragraph(meeting.summary)
    document.add_heading("Поручения", 1)
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Поручение", "Ответственный", "Срок"]): cell.text = value
    for task in meeting.tasks:
        cells = table.add_row().cells
        for cell, value in zip(cells, [task.task, task.assignee or "Уточнить", task.deadline or "Не указан"]): cell.text = value
    document.add_heading("Полный транскрипт", 1)
    for segment in (meeting.transcript or {}).get("segments", []):
        speaker = meeting.speaker_mapping.get(segment["speaker"], segment["speaker"])
        paragraph = document.add_paragraph(f"[{_time(segment['start'])}] {speaker}: {segment['text']}")
        paragraph.style.font.size = Pt(9)
    output = BytesIO(); document.save(output); return output.getvalue()


def pdf_bytes(meeting) -> bytes:
    output = BytesIO(); styles = getSampleStyleSheet()
    font_name = "Helvetica"
    for candidate in (Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")):
        if candidate.exists():
            font_name = "ProtocolUnicode"
            if font_name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font_name, str(candidate)))
            break
    for style in styles.byName.values():
        style.fontName = font_name
    content = [Paragraph("ПРОТОКОЛ СОВЕЩАНИЯ", styles["Title"]), Spacer(1, 12), Paragraph(f"Дата: {meeting.meeting_date}", styles["BodyText"]), Paragraph(meeting.summary, styles["BodyText"]), Spacer(1, 12)]
    rows = [["Поручение", "Ответственный", "Срок"]] + [[task.task, task.assignee or "Уточнить", task.deadline or "Не указан"] for task in meeting.tasks]
    table = Table(rows, colWidths=[260, 130, 100])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, -1), font_name), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    content.extend([Paragraph("Поручения", styles["Heading2"]), table, Spacer(1, 12), Paragraph("Полный транскрипт", styles["Heading2"])])
    for segment in (meeting.transcript or {}).get("segments", []):
        speaker = meeting.speaker_mapping.get(segment["speaker"], segment["speaker"])
        content.append(Paragraph(f"[{_time(segment['start'])}] {speaker}: {segment['text']}", styles["BodyText"]))
    SimpleDocTemplate(output, pagesize=A4, rightMargin=36, leftMargin=36).build(content)
    return output.getvalue()
